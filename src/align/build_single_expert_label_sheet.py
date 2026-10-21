from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path("Paper 4")
VALID_LABELS = {"MATCH", "PARTIAL", "NO_MATCH"}
SHEET_NAME = "single_expert_labels"
LABEL_COLUMNS = ["human_label", "human_notes", "expert_confidence"]


def norm_label(value: object) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip().upper()
    return text if text in VALID_LABELS else ""


def load_source(seed_path: Path, diagnostics_path: Path, priority_path: Path) -> pd.DataFrame:
    seed = pd.read_csv(seed_path, low_memory=False)
    source = seed.copy()

    if diagnostics_path.exists():
        diagnostics = pd.read_csv(diagnostics_path, low_memory=False)
        keep = [
            "candidate_rank",
            "ambiguity_score",
            "diagnostic_reason",
            "ai_suggested_label",
            "ai_confidence",
            "lexical_overlap_value",
            "event_count_ratio",
        ]
        diagnostics = diagnostics[[col for col in keep if col in diagnostics.columns]].drop_duplicates("candidate_rank")
        source = source.merge(diagnostics, on="candidate_rank", how="left")
    else:
        for col in ["ambiguity_score", "diagnostic_reason", "ai_suggested_label", "ai_confidence", "lexical_overlap_value", "event_count_ratio"]:
            source[col] = ""

    if priority_path.exists():
        priority = pd.read_csv(priority_path, low_memory=False)
        keep = ["candidate_rank", "review_priority_score", "review_reason", "llm_rationale"]
        priority = priority[[col for col in keep if col in priority.columns]].drop_duplicates("candidate_rank")
        source = source.merge(priority, on="candidate_rank", how="left")
    else:
        for col in ["review_priority_score", "review_reason", "llm_rationale"]:
            source[col] = ""

    for col in LABEL_COLUMNS:
        if col not in source.columns:
            source[col] = ""
    source["human_label"] = source["human_label"].map(norm_label)
    source["human_notes"] = source["human_notes"].fillna("").astype(str)
    source["expert_confidence"] = source["expert_confidence"].fillna("").astype(str)
    return source


def merge_existing_sheet(df: pd.DataFrame, workbook_path: Path) -> tuple[pd.DataFrame, dict]:
    result = {
        "existing_workbook_found": workbook_path.exists(),
        "preserved_rows": 0,
        "preserved_cells": 0,
    }
    if not workbook_path.exists():
        return df, result

    try:
        existing = pd.read_excel(workbook_path, sheet_name=SHEET_NAME)
    except ValueError:
        return df, result
    if "candidate_rank" not in existing.columns:
        return df, result

    keep = ["candidate_rank", *[col for col in LABEL_COLUMNS if col in existing.columns]]
    existing = existing[keep].copy()
    existing["candidate_rank"] = pd.to_numeric(existing["candidate_rank"], errors="coerce").astype("Int64")
    existing = existing.dropna(subset=["candidate_rank"]).drop_duplicates("candidate_rank", keep="last")
    if "human_label" not in existing.columns:
        existing["human_label"] = ""
    existing["human_label_norm"] = existing["human_label"].map(norm_label)
    invalid = existing[
        existing["human_label"].notna()
        & existing["human_label"].astype(str).str.strip().ne("")
        & existing["human_label_norm"].eq("")
    ]
    if not invalid.empty:
        ranks = invalid["candidate_rank"].dropna().astype(int).tolist()
        raise ValueError(f"Invalid labels found in existing single-expert workbook at candidate_rank values: {ranks}")
    existing["human_label"] = existing["human_label_norm"].fillna("")
    for col in ["human_notes", "expert_confidence"]:
        if col not in existing.columns:
            existing[col] = ""
        existing[col] = existing[col].fillna("").astype(str)

    merged = df.merge(
        existing[["candidate_rank", *LABEL_COLUMNS]],
        on="candidate_rank",
        how="left",
        suffixes=("", "_existing"),
    )
    preserved_rows: set[int] = set()
    preserved_cells = 0
    for col in LABEL_COLUMNS:
        existing_col = f"{col}_existing"
        if existing_col not in merged.columns:
            continue
        nonblank = merged[existing_col].fillna("").astype(str).str.strip().ne("")
        if nonblank.any():
            merged.loc[nonblank, col] = merged.loc[nonblank, existing_col]
            preserved_cells += int(nonblank.sum())
            preserved_rows.update(merged.loc[nonblank, "candidate_rank"].astype(int).tolist())
        merged = merged.drop(columns=[existing_col])
    result["preserved_rows"] = len(preserved_rows)
    result["preserved_cells"] = preserved_cells
    return merged, result


def prepare_sheet(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for col in ["ambiguity_score", "review_priority_score", "ai_confidence", "lexical_overlap_value", "event_count_ratio"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    df["is_labeled"] = df["human_label"].map(norm_label).ne("")
    df = df.sort_values(
        ["is_labeled", "ambiguity_score", "review_priority_score", "candidate_rank"],
        ascending=[True, False, False, True],
    ).reset_index(drop=True)
    df["review_order"] = range(1, len(df) + 1)
    df["priority_group"] = df["review_order"].map(lambda value: "first_pass_30" if int(value) <= 30 else "full_pass")
    df["label_status"] = df["human_label"].map(lambda value: "labeled" if norm_label(value) else "needs_label")

    keep = [
        "review_order",
        "priority_group",
        "label_status",
        "candidate_rank",
        "human_label",
        "human_notes",
        "expert_confidence",
        "taxonomy_code",
        "dart_event",
        "dart_subtype",
        "sec_event",
        "sec_subtype",
        "proposed_label",
        "ai_suggested_label",
        "ai_confidence",
        "ambiguity_score",
        "diagnostic_reason",
        "review_priority_score",
        "review_reason",
        "lexical_overlap_value",
        "event_count_ratio",
        "llm_rationale",
        "rationale",
    ]
    for col in keep:
        if col not in df.columns:
            df[col] = ""
    return df[keep]


def write_workbook(df: pd.DataFrame, output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    headers = list(df.columns)
    ws.append(headers)
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "A": 12,
        "B": 14,
        "C": 14,
        "D": 14,
        "E": 16,
        "F": 42,
        "G": 18,
        "H": 14,
        "I": 34,
        "J": 28,
        "K": 44,
        "L": 28,
        "M": 16,
        "N": 18,
        "O": 14,
        "P": 14,
        "Q": 52,
        "R": 18,
        "S": 44,
        "T": 18,
        "U": 18,
        "V": 60,
        "W": 48,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    label_validation = DataValidation(type="list", formula1='"MATCH,PARTIAL,NO_MATCH"', allow_blank=True)
    ws.add_data_validation(label_validation)
    label_validation.add(f"E2:E{len(df) + 1}")

    confidence_validation = DataValidation(type="list", formula1='"high,medium,low"', allow_blank=True)
    ws.add_data_validation(confidence_validation)
    confidence_validation.add(f"G2:G{len(df) + 1}")

    guide = wb.create_sheet("guide")
    guide_rows = [
        ["Purpose", "Fast single-expert labeling route for RQ1 when a two-annotator workflow is not available."],
        ["MATCH", "The DART and SEC events describe substantially the same legal/economic disclosure event."],
        ["PARTIAL", "The events share a risk domain but differ in legal trigger, timing, or specificity."],
        ["NO_MATCH", "The relationship is too broad or unrelated."],
        ["Workflow", 'Fill human_label in the single_expert_labels sheet, then run: python "Paper 4\\src\\align\\sync_single_expert_label_sheet.py"'],
        ["Guardrail", "Single-expert labels can support Precision@1/MRR with a single-expert-gold caveat, but cannot support inter-annotator agreement or Cohen's kappa."],
    ]
    for row in guide_rows:
        guide.append(row)
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 110
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in guide["A"]:
        cell.font = Font(bold=True)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def workbook_structure_check(workbook_path: Path, expected_rows: int, expected_columns: list[str]) -> dict:
    checks: list[dict] = []
    wb = load_workbook(workbook_path, data_only=False)

    def add_check(name: str, ok: bool, detail: str = "") -> None:
        checks.append({"check": name, "status": "pass" if ok else "fail", "detail": detail})

    expected_sheets = {SHEET_NAME, "guide"}
    add_check("required_sheets", expected_sheets.issubset(set(wb.sheetnames)), ",".join(wb.sheetnames))
    if SHEET_NAME in wb.sheetnames:
        ws = wb[SHEET_NAME]
        headers = [cell.value for cell in ws[1]]
        add_check("label_sheet_row_count", ws.max_row == expected_rows + 1, f"{ws.max_row}/{expected_rows + 1}")
        add_check("label_sheet_columns", headers == expected_columns, f"{len(headers)}/{len(expected_columns)}")
        validation_refs = ";".join(str(validation.sqref) for validation in ws.data_validations.dataValidation)
        add_check("human_label_dropdown_range", f"E2:E{expected_rows + 1}" in validation_refs, validation_refs)
        add_check("confidence_dropdown_range", f"G2:G{expected_rows + 1}" in validation_refs, validation_refs)
        add_check("frozen_header", ws.freeze_panes == "A2", str(ws.freeze_panes))
    if "guide" in wb.sheetnames:
        guide_text = "\n".join(str(cell.value or "") for row in wb["guide"].iter_rows() for cell in row)
        add_check("guide_single_expert_guardrail", "cannot support inter-annotator agreement" in guide_text, "")
    return {
        "status": "pass" if all(row["status"] == "pass" for row in checks) else "fail",
        "checks": checks,
    }


def write_guide(summary: dict, path: Path) -> None:
    lines = [
        "# Single-Expert Alignment Labeling Guide",
        "",
        "This is the fast RQ1 labeling route when a two-annotator workflow is not feasible. It uses the existing `human_label` fallback that `evaluate_alignment.py` already supports.",
        "",
        "## Files",
        "",
        f"- Excel sheet: `{summary['xlsx_path']}`",
        f"- CSV mirror: `{summary['csv_path']}`",
        "",
        "## Label Set",
        "",
        "| Label | Use when |",
        "|---|---|",
        "| `MATCH` | The DART and SEC events describe substantially the same legal/economic disclosure event. |",
        "| `PARTIAL` | The events share a risk domain but differ in legal trigger, timing, or specificity. |",
        "| `NO_MATCH` | The relationship is too broad or unrelated. |",
        "",
        "## Workflow",
        "",
        "1. Open the Excel sheet and fill `human_label` for all 100 rows.",
        "2. Use `human_notes` for borderline cases and `expert_confidence` for `high`, `medium`, or `low` confidence.",
        "3. Run `python \"Paper 4\\src\\align\\sync_single_expert_label_sheet.py\"`.",
        "4. Rerun `python \"Paper 4\\src\\align\\evaluate_alignment.py\"` if it was not run by the sync command.",
        "",
        "## Guardrail",
        "",
        "After `human_label` is filled and synced, single-expert labels can be reported as a single-expert gold evaluation for Precision@1/MRR. They cannot support inter-annotator agreement or Cohen's kappa.",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_summary(summary: dict, path: Path) -> None:
    lines = [
        "# Single-Expert Label Sheet Readiness",
        "",
        "| Item | Value |",
        "|---|---:|",
        f"| Candidate rows | {summary['candidate_rows']} |",
        f"| Human-labeled rows | {summary['human_labeled_rows']} |",
        f"| Unlabeled rows | {summary['unlabeled_rows']} |",
        f"| First-pass rows | {summary['first_pass_rows']} |",
        f"| High-ambiguity rows | {summary['high_ambiguity_rows']} |",
        f"| Preserved existing rows | {summary['preservation']['preserved_rows']} |",
        f"| Preserved existing cells | {summary['preservation']['preserved_cells']} |",
        f"| Workbook structure status | {summary['workbook_structure']['status']} |",
        "",
        "## Guardrail",
        "",
        summary["guardrail"],
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def build_sheet(
    seed_path: Path,
    diagnostics_path: Path,
    priority_path: Path,
    out_csv: Path,
    out_xlsx: Path,
    out_json: Path,
    out_md: Path,
    guide_path: Path,
) -> dict:
    df = load_source(seed_path, diagnostics_path, priority_path)
    df, preservation = merge_existing_sheet(df, out_xlsx)
    sheet = prepare_sheet(df)
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    sheet.to_csv(out_csv, index=False, encoding="utf-8-sig")
    write_workbook(sheet, out_xlsx)
    human_labeled_rows = int(sheet["human_label"].map(norm_label).ne("").sum())
    high_ambiguity_rows = int(pd.to_numeric(sheet["ambiguity_score"], errors="coerce").fillna(0).ge(3).sum())
    summary = {
        "status": "ready_for_single_expert_labeling" if human_labeled_rows == 0 else "partially_labeled",
        "candidate_rows": int(len(sheet)),
        "human_labeled_rows": human_labeled_rows,
        "unlabeled_rows": int(len(sheet) - human_labeled_rows),
        "first_pass_rows": int(sheet["priority_group"].eq("first_pass_30").sum()),
        "high_ambiguity_rows": high_ambiguity_rows,
        "csv_path": str(out_csv),
        "xlsx_path": str(out_xlsx),
        "guide_path": str(guide_path),
        "preservation": preservation,
        "workbook_structure": workbook_structure_check(out_xlsx, len(sheet), list(sheet.columns)),
        "guardrail": "This sheet supports single-expert RQ1 labeling only; do not report inter-annotator agreement from it.",
    }
    if summary["workbook_structure"]["status"] != "pass":
        summary["status"] = "workbook_structure_failed"
    out_json.parent.mkdir(parents=True, exist_ok=True)
    out_json.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    write_summary(summary, out_md)
    write_guide(summary, guide_path)
    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--seed-path", default=str(ROOT / "data" / "processed" / "manual_alignment_seed_100.csv"))
    parser.add_argument("--diagnostics-path", default=str(ROOT / "data" / "processed" / "manual_alignment_seed_100_candidate_diagnostics.csv"))
    parser.add_argument("--priority-path", default=str(ROOT / "data" / "processed" / "manual_alignment_seed_100_annotation_priority.csv"))
    parser.add_argument("--out-csv", default=str(ROOT / "data" / "processed" / "manual_alignment_seed_100_single_expert_label_sheet.csv"))
    parser.add_argument("--out-xlsx", default=str(ROOT / "data" / "processed" / "manual_alignment_seed_100_single_expert_label_sheet.xlsx"))
    parser.add_argument("--out-json", default=str(ROOT / "results" / "alignment_single_expert_readiness.json"))
    parser.add_argument("--out-md", default=str(ROOT / "results" / "alignment_single_expert_readiness.md"))
    parser.add_argument("--guide-path", default=str(ROOT / "manuscript" / "alignment_single_expert_labeling_guide.md"))
    args = parser.parse_args()
    result = build_sheet(
        Path(args.seed_path),
        Path(args.diagnostics_path),
        Path(args.priority_path),
        Path(args.out_csv),
        Path(args.out_xlsx),
        Path(args.out_json),
        Path(args.out_md),
        Path(args.guide_path),
    )
    print(json.dumps(result, indent=2, ensure_ascii=False))
    if result["workbook_structure"]["status"] != "pass":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
