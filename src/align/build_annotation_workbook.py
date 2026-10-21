from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from evaluate_alignment import normalize_label


VALID_LABELS = ["MATCH", "PARTIAL", "NO_MATCH"]
LABEL_COLUMNS = [
    "annotator1_label",
    "annotator2_label",
    "adjudicated_label",
    "human_label",
]
PRESERVED_COLUMNS = [
    *LABEL_COLUMNS,
    "human_notes",
    "adjudication_notes",
]
ANNOTATION_COLUMNS = [
    "candidate_rank",
    "taxonomy_code",
    "dart_event",
    "dart_subtype",
    "sec_event",
    "sec_subtype",
    "proposed_label",
    "ai_suggested_label",
    "ai_confidence",
    "ai_rationale",
    "review_priority_score",
    "review_reason",
    "annotator1_label",
    "annotator2_label",
    "adjudicated_label",
    "human_label",
    "human_notes",
    "adjudication_notes",
    "rationale",
    "dart_event_count",
    "sec_event_count",
]


def load_annotation(seed_path: Path, llm_path: Path | None) -> pd.DataFrame:
    seed = pd.read_csv(seed_path, low_memory=False).copy()
    if llm_path and llm_path.exists():
        llm = pd.read_csv(llm_path, low_memory=False)
        keep = [
            "candidate_rank",
            "llm_decision",
            "llm_confidence",
            "llm_rationale",
            "llm_mode",
            "api_status",
        ]
        llm = llm[[col for col in keep if col in llm.columns]].drop_duplicates("candidate_rank")
        seed = seed.merge(llm, on="candidate_rank", how="left")
    else:
        seed["llm_decision"] = ""
        seed["llm_confidence"] = ""
        seed["llm_rationale"] = ""
        seed["llm_mode"] = ""
        seed["api_status"] = ""

    seed["ai_suggested_label"] = seed.get("llm_decision", "").fillna("")
    seed["ai_confidence"] = seed.get("llm_confidence", "").fillna("")
    seed["ai_rationale"] = seed.get("llm_rationale", "").fillna("")
    proposed = seed.get("proposed_label", "").fillna("").astype(str).str.strip().str.upper()
    ai = seed["ai_suggested_label"].fillna("").astype(str).str.strip().str.upper()
    conf = pd.to_numeric(seed["ai_confidence"], errors="coerce")
    reasons = []
    scores = []
    for idx in seed.index:
        row_reasons = []
        score = 50
        if proposed.loc[idx] != ai.loc[idx]:
            score += 30
            row_reasons.append("proposed_vs_ai_disagreement")
        if pd.isna(conf.loc[idx]) or conf.loc[idx] < 0.65:
            score += 20
            row_reasons.append("low_or_missing_ai_confidence")
        if str(seed.loc[idx].get("taxonomy_code", "")) in {"T05", "T06"}:
            score += 5
            row_reasons.append("broad_taxonomy_cell")
        if not row_reasons:
            row_reasons.append("routine_confirmation")
        scores.append(score)
        reasons.append("; ".join(row_reasons))
    seed["review_priority_score"] = scores
    seed["review_reason"] = reasons
    for col in ANNOTATION_COLUMNS:
        if col not in seed.columns:
            seed[col] = ""
    return seed[ANNOTATION_COLUMNS].sort_values("candidate_rank")


def merge_existing_annotations(df: pd.DataFrame, workbook_path: Path) -> tuple[pd.DataFrame, dict]:
    result = {
        "existing_workbook_found": workbook_path.exists(),
        "preserved_rows": 0,
        "preserved_cells": 0,
        "preserved_columns": PRESERVED_COLUMNS,
    }
    if not workbook_path.exists():
        return df, result

    existing = pd.read_excel(workbook_path, sheet_name="annotation")
    if "candidate_rank" not in existing.columns:
        raise ValueError(f"Existing workbook is missing candidate_rank: {workbook_path}")

    keep = ["candidate_rank", *[col for col in PRESERVED_COLUMNS if col in existing.columns]]
    existing = existing[keep].copy()
    existing["candidate_rank"] = pd.to_numeric(existing["candidate_rank"], errors="coerce").astype("Int64")
    existing = existing.dropna(subset=["candidate_rank"]).drop_duplicates("candidate_rank", keep="last")

    invalid_ranks: set[int] = set()
    for col in LABEL_COLUMNS:
        if col not in existing.columns:
            existing[col] = ""
        existing[f"{col}_norm"] = existing[col].map(normalize_label)
        nonblank = existing[col].notna() & existing[col].astype(str).str.strip().ne("")
        invalid = existing[nonblank & existing[f"{col}_norm"].eq("")]
        invalid_ranks.update(invalid["candidate_rank"].astype(int).tolist())
        existing[col] = existing[f"{col}_norm"].fillna("")
    if invalid_ranks:
        raise ValueError(f"Invalid labels found in existing workbook at candidate_rank values: {sorted(invalid_ranks)}")

    for col in PRESERVED_COLUMNS:
        if col not in existing.columns:
            existing[col] = ""
        existing[col] = existing[col].fillna("").astype(str)

    merged = df.merge(
        existing[["candidate_rank", *PRESERVED_COLUMNS]],
        on="candidate_rank",
        how="left",
        suffixes=("", "_existing"),
    )
    preserved_rows: set[int] = set()
    preserved_cells = 0
    for col in PRESERVED_COLUMNS:
        existing_col = f"{col}_existing"
        if existing_col not in merged.columns:
            continue
        existing_nonblank = merged[existing_col].fillna("").astype(str).str.strip().ne("")
        if existing_nonblank.any():
            merged.loc[existing_nonblank, col] = merged.loc[existing_nonblank, existing_col]
            preserved_cells += int(existing_nonblank.sum())
            preserved_rows.update(merged.loc[existing_nonblank, "candidate_rank"].astype(int).tolist())
        merged = merged.drop(columns=[existing_col])

    result["preserved_rows"] = len(preserved_rows)
    result["preserved_cells"] = preserved_cells
    return merged[ANNOTATION_COLUMNS].sort_values("candidate_rank"), result


def write_workbook(df: pd.DataFrame, output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "annotation"

    headers = list(df.columns)
    ws.append(headers)
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "A": 14,
        "B": 14,
        "C": 32,
        "D": 28,
        "E": 42,
        "F": 28,
        "G": 16,
        "H": 18,
        "I": 14,
        "J": 54,
        "K": 16,
        "L": 42,
        "M": 18,
        "N": 18,
        "O": 18,
        "P": 16,
        "Q": 36,
        "R": 42,
        "S": 48,
        "T": 16,
        "U": 16,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    dv = DataValidation(type="list", formula1='"MATCH,PARTIAL,NO_MATCH"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"M2:P{len(df) + 1}")

    mismatch_fill = PatternFill("solid", fgColor="FFF2CC")
    ws.conditional_formatting.add(
        f"O2:P{len(df) + 1}",
        FormulaRule(formula=[f'AND(O2<>"",O2<>$H2)'], fill=mismatch_fill),
    )
    disagreement_fill = PatternFill("solid", fgColor="F4CCCC")
    ws.conditional_formatting.add(
        f"M2:N{len(df) + 1}",
        FormulaRule(formula=[f'AND($M2<>"",$N2<>"",$M2<>$N2)'], fill=disagreement_fill),
    )

    guide = wb.create_sheet("guide")
    guide_rows = [
        ["Purpose", "Fill annotator1_label and annotator2_label independently. Use adjudicated_label after resolving disagreements. Do not edit candidate_rank."],
        ["MATCH", "The DART and SEC events describe substantially the same legal/economic disclosure event."],
        ["PARTIAL", "The events share a risk theme but differ in legal trigger, timing, or specificity."],
        ["NO_MATCH", "The relationship is too broad or unrelated."],
        ["Workflow", 'After labeling, run: python "Paper 4\\src\\align\\sync_annotation_workbook.py"'],
        ["Gold label", "evaluate_alignment.py uses adjudicated_label first, then legacy human_label, then matching annotator consensus. Disagreements without adjudication are not scored."],
        ["Guardrail", "AI suggested labels are review aids only and must not be reported as human labels."],
    ]
    for row in guide_rows:
        guide.append(row)
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 96
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in guide["A"]:
        cell.font = Font(bold=True)

    summary = wb.create_sheet("summary")
    summary_rows = [
        ["Metric", "Value"],
        ["Candidate rows", len(df)],
        ["Annotator 1 labeled rows", '=COUNTIF(annotation!M:M,"MATCH")+COUNTIF(annotation!M:M,"PARTIAL")+COUNTIF(annotation!M:M,"NO_MATCH")'],
        ["Annotator 2 labeled rows", '=COUNTIF(annotation!N:N,"MATCH")+COUNTIF(annotation!N:N,"PARTIAL")+COUNTIF(annotation!N:N,"NO_MATCH")'],
        ["Adjudicated rows", '=COUNTIF(annotation!O:O,"MATCH")+COUNTIF(annotation!O:O,"PARTIAL")+COUNTIF(annotation!O:O,"NO_MATCH")'],
        ["Legacy human_label rows", '=COUNTIF(annotation!P:P,"MATCH")+COUNTIF(annotation!P:P,"PARTIAL")+COUNTIF(annotation!P:P,"NO_MATCH")'],
        ["Annotator exact agreements", '=SUMPRODUCT(--(annotation!M2:M101<>""),--(annotation!M2:M101=annotation!N2:N101))'],
        ["Annotator disagreements", '=SUMPRODUCT(--(annotation!M2:M101<>""),--(annotation!N2:N101<>""),--(annotation!M2:M101<>annotation!N2:N101))'],
        ["Blank adjudicated labels", '=COUNTBLANK(annotation!O2:O101)'],
    ]
    for row in summary_rows:
        summary.append(row)
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 18
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def workbook_structure_check(workbook_path: Path, expected_rows: int) -> dict:
    checks: list[dict] = []
    wb = load_workbook(workbook_path, data_only=False)

    def add_check(name: str, ok: bool, detail: str = "") -> None:
        checks.append({"check": name, "status": "pass" if ok else "fail", "detail": detail})

    expected_sheets = {"annotation", "guide", "summary"}
    add_check("required_sheets", expected_sheets.issubset(set(wb.sheetnames)), ",".join(wb.sheetnames))
    if "annotation" in wb.sheetnames:
        ws = wb["annotation"]
        headers = [cell.value for cell in ws[1]]
        add_check("annotation_row_count", ws.max_row == expected_rows + 1, f"{ws.max_row}/{expected_rows + 1}")
        add_check("annotation_columns", headers == ANNOTATION_COLUMNS, f"{len(headers)}/{len(ANNOTATION_COLUMNS)}")
        validation_refs = ";".join(str(validation.sqref) for validation in ws.data_validations.dataValidation)
        add_check("label_dropdown_range", f"M2:P{expected_rows + 1}" in validation_refs, validation_refs)
        add_check("frozen_header", ws.freeze_panes == "A2", str(ws.freeze_panes))
    if "guide" in wb.sheetnames:
        guide_text = "\n".join(str(cell.value or "") for row in wb["guide"].iter_rows() for cell in row)
        add_check("guide_guardrail", "AI suggested labels are review aids only" in guide_text, "")
    if "summary" in wb.sheetnames:
        summary = wb["summary"]
        formulas = [str(cell.value or "") for row in summary.iter_rows() for cell in row if str(cell.value or "").startswith("=")]
        add_check("summary_formulas", len(formulas) >= 7, str(len(formulas)))
    return {
        "status": "pass" if all(row["status"] == "pass" for row in checks) else "fail",
        "checks": checks,
    }


def write_readiness(df: pd.DataFrame, output_json: Path, output_md: Path, workbook_path: Path, preservation: dict | None = None) -> dict:
    for col in ["annotator1_label", "annotator2_label", "adjudicated_label", "human_label"]:
        if col not in df.columns:
            df[col] = ""
        df[f"{col}_norm"] = df[col].map(normalize_label)
    gold = []
    source = []
    for row in df.itertuples(index=False):
        adjudicated = getattr(row, "adjudicated_label_norm")
        legacy = getattr(row, "human_label_norm")
        ann1 = getattr(row, "annotator1_label_norm")
        ann2 = getattr(row, "annotator2_label_norm")
        if adjudicated:
            gold.append(adjudicated)
            source.append("adjudicated_label")
        elif legacy:
            gold.append(legacy)
            source.append("human_label")
        elif ann1 and ann2 and ann1 == ann2:
            gold.append(ann1)
            source.append("annotator_consensus")
        elif ann1 and ann2:
            gold.append("")
            source.append("needs_adjudication")
        elif ann1:
            gold.append("")
            source.append("annotator1_only")
        elif ann2:
            gold.append("")
            source.append("annotator2_only")
        else:
            gold.append("")
            source.append("none")
    df["gold_label_norm"] = gold
    df["gold_label_source"] = source
    result = {
        "candidate_rows": int(len(df)),
        "human_labeled_rows": int(df["gold_label_norm"].isin(VALID_LABELS).sum()),
        "human_unlabeled_rows": int((~df["gold_label_norm"].isin(VALID_LABELS)).sum()),
        "annotator1_labeled_rows": int(df["annotator1_label_norm"].isin(VALID_LABELS).sum()),
        "annotator2_labeled_rows": int(df["annotator2_label_norm"].isin(VALID_LABELS).sum()),
        "adjudicated_labeled_rows": int(df["adjudicated_label_norm"].isin(VALID_LABELS).sum()),
        "needs_adjudication_rows": int(df["gold_label_source"].eq("needs_adjudication").sum()),
        "taxonomy_counts": {str(k): int(v) for k, v in df["taxonomy_code"].value_counts().items()},
        "proposed_label_counts": {str(k): int(v) for k, v in df["proposed_label"].value_counts().items()},
        "ai_suggested_label_counts": {str(k): int(v) for k, v in df["ai_suggested_label"].value_counts().items()},
        "workbook": str(workbook_path),
        "preservation": preservation or {},
        "workbook_structure": workbook_structure_check(workbook_path, len(df)),
        "status": "ready_for_human_labeling" if int(df["gold_label_norm"].isin(VALID_LABELS).sum()) == 0 else "partially_labeled",
    }
    if result["workbook_structure"]["status"] != "pass":
        result["status"] = "workbook_structure_failed"
    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_json.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")

    lines = [
        "# Alignment Annotation Readiness",
        "",
        "| Item | Value |",
        "|---|---:|",
        f"| Candidate rows | {result['candidate_rows']} |",
        f"| Gold-labeled rows | {result['human_labeled_rows']} |",
        f"| Unlabeled rows | {result['human_unlabeled_rows']} |",
        f"| Annotator 1 labeled rows | {result['annotator1_labeled_rows']} |",
        f"| Annotator 2 labeled rows | {result['annotator2_labeled_rows']} |",
        f"| Adjudicated labeled rows | {result['adjudicated_labeled_rows']} |",
        f"| Rows needing adjudication | {result['needs_adjudication_rows']} |",
        f"| Preserved existing annotation rows | {(preservation or {}).get('preserved_rows', 0)} |",
        f"| Preserved existing annotation cells | {(preservation or {}).get('preserved_cells', 0)} |",
        f"| Workbook structure status | {result['workbook_structure']['status']} |",
        f"| Workbook | `{result['workbook']}` |",
        "",
        "## Guardrail",
        "",
        "AI suggested labels are review aids only. Do not report Precision@1 or MRR until `annotator1_label`/`annotator2_label` are filled, disagreements are resolved in `adjudicated_label`, and `evaluate_alignment.py` is rerun.",
    ]
    output_md.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--seed-path", default=str(Path("Paper 4") / "data" / "processed" / "manual_alignment_seed_100.csv"))
    parser.add_argument("--llm-path", default=str(Path("Paper 4") / "data" / "processed" / "llm_alignment_results.csv"))
    parser.add_argument("--output-path", default=str(Path("Paper 4") / "data" / "processed" / "manual_alignment_seed_100_annotation_workbook.xlsx"))
    parser.add_argument("--out-json", default=str(Path("Paper 4") / "results" / "alignment_annotation_readiness.json"))
    parser.add_argument("--out-md", default=str(Path("Paper 4") / "results" / "alignment_annotation_readiness.md"))
    args = parser.parse_args()

    output_path = Path(args.output_path)
    df = load_annotation(Path(args.seed_path), Path(args.llm_path) if args.llm_path else None)
    df, preservation = merge_existing_annotations(df, output_path)
    write_workbook(df, output_path)
    result = write_readiness(df, Path(args.out_json), Path(args.out_md), output_path, preservation)
    print(json.dumps(result, indent=2, ensure_ascii=False))
    if result["workbook_structure"]["status"] != "pass":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
